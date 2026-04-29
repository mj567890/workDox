from fastapi import APIRouter, Depends, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, or_
from sqlalchemy.orm import selectinload
import asyncio
from datetime import datetime, timezone, timedelta
from io import BytesIO

from app.dependencies import get_current_user, get_db, check_permission
from app.core.exceptions import NotFoundException, ConflictException, ForbiddenException, ValidationException
from app.core.permissions import Permission, RoleCode
from app.core.pagination import PaginationParams
from app.core.storage import minio_client
from app.models.user import User
from app.models.document import (
    Document, DocumentVersion, DocumentCategory, Tag,
    DocumentEditLock, DocumentReview, CrossMatterReference,
)
from app.models.matter import Matter, MatterMember
from app.utils.file_utils import detect_file_type, generate_storage_path, compute_sha256, is_allowed_file
from app.services.document_service import (
    DocumentService, DocumentVersionService, DocumentLockService, CrossReferenceService, DocumentReviewService,
)
from app.services.document_intelligence import extract_and_store_text, find_similar_by_vector
from app.utils.text_extractor import extract_text, is_supported
from openpyxl import Workbook

router = APIRouter()
document_service = DocumentService()
version_service = DocumentVersionService()
lock_service = DocumentLockService()
reference_service = CrossReferenceService()
review_service = DocumentReviewService()


# ---------- Pydantic Schemas ----------

class DocumentCreate(BaseModel):
    original_name: str
    file_type: str
    file_size: int
    mime_type: str
    storage_path: str
    description: str | None = None
    matter_id: int | None = None
    category_id: int | None = None
    tag_ids: list[int] = []
    checksum: str | None = None


class DocumentUpdate(BaseModel):
    original_name: str | None = None
    description: str | None = None
    category_id: int | None = None
    tag_ids: list[int] | None = None
    permission_scope: str | None = None


class DocumentOut(BaseModel):
    id: int
    original_name: str
    file_type: str
    file_size: int
    mime_type: str
    description: str | None
    owner_id: int
    owner_name: str | None
    matter_id: int | None
    matter_title: str | None
    category_id: int | None
    category_name: str | None
    status: str
    current_version_id: int | None
    current_version_no: int | None
    permission_scope: str
    is_deleted: bool
    tags: list[str]
    created_at: str | None
    updated_at: str | None
    extracted_text_length: int | None = None

    class Config:
        from_attributes = True


class DocumentListResponse(BaseModel):
    items: list[DocumentOut]
    total: int
    page: int
    page_size: int
    total_pages: int


class VersionOut(BaseModel):
    id: int
    version_no: int
    file_size: int
    upload_user_name: str | None
    change_note: str | None
    is_official: bool
    checksum: str | None
    created_at: str | None

    class Config:
        from_attributes = True


class CategoryCreate(BaseModel):
    name: str
    code: str
    description: str | None = None
    sort_order: int = 0


class CategoryUpdate(BaseModel):
    name: str | None = None
    description: str | None = None
    sort_order: int | None = None


class CategoryOut(BaseModel):
    id: int
    name: str
    code: str
    description: str | None
    sort_order: int
    is_system: bool

    class Config:
        from_attributes = True


class TagCreate(BaseModel):
    name: str
    color: str = "#409EFF"


class TagOut(BaseModel):
    id: int
    name: str
    color: str

    class Config:
        from_attributes = True


class ReferenceCreate(BaseModel):
    matter_id: int
    is_readonly: bool = True


class ReferenceOut(BaseModel):
    id: int
    matter_id: int
    matter_title: str | None
    is_readonly: bool
    added_by_name: str | None
    created_at: str | None

    class Config:
        from_attributes = True


class LockStatusOut(BaseModel):
    is_locked: bool
    locked_by: int | None
    locked_by_name: str | None
    locked_at: str | None
    expires_at: str | None


class PendingUpload(BaseModel):
    id: str
    filename: str
    uploaded_chunks: int
    total_chunks: int
    is_complete: bool


class CompleteUploadRequest(BaseModel):
    upload_id: str
    filename: str
    matter_id: int | None = None
    category_id: int | None = None
    description: str | None = None
    tag_ids: list[int] = []


# ---------- Helper Functions ----------

async def _background_extract(db: AsyncSession, doc_id: int, file_data: bytes, file_type: str, original_name: str):
    """Background task to extract text from uploaded document."""
    try:
        extracted = await extract_and_store_text(db, doc_id, file_data, file_type, original_name)
        if extracted:
            # Dispatch Celery task for embedding (uses independent session)
            from app.tasks.embedding_tasks import embed_document_task
            embed_document_task.delay(doc_id)
    except Exception:
        pass  # Silently ignore extraction failures in background


def _user_to_dict(user: User) -> dict:
    """Convert User ORM object to a dict for service current_user param."""
    role_codes = [r.role_code for r in (user.roles or []) if hasattr(r, 'role_code')]
    return {"id": user.id, "roles": role_codes}


async def _check_document_access(user: User, doc: Document, db: AsyncSession):
    """Raise ForbiddenException if user cannot access document."""
    role_codes = [r.role_code for r in (user.roles or []) if hasattr(r, 'role_code')]
    if RoleCode.ADMIN in role_codes:
        return

    if doc.owner_id == user.id:
        return

    if doc.matter_id:
        member_result = await db.execute(
            select(MatterMember).where(
                MatterMember.matter_id == doc.matter_id,
                MatterMember.user_id == user.id,
            )
        )
        if member_result.scalar_one_or_none():
            return

        if RoleCode.DEPT_LEADER in role_codes:
            matter_result = await db.execute(
                select(Matter.owner_id).where(Matter.id == doc.matter_id)
            )
            matter_owner_id = matter_result.scalar_one_or_none()
            if matter_owner_id:
                owner_result = await db.execute(
                    select(User.department_id).where(User.id == matter_owner_id)
                )
                matter_dept_id = owner_result.scalar_one_or_none()
                if matter_dept_id and user.department_id == matter_dept_id:
                    return

    raise ForbiddenException(detail="Access denied to this document")


def _doc_to_out(d: Document) -> DocumentOut:
    """Format a Document ORM object to DocumentOut response schema."""
    current_version_no = None
    if d.current_version:
        current_version_no = d.current_version.version_no
    elif d.current_version_id and d.versions:
        for v in d.versions:
            if v.id == d.current_version_id:
                current_version_no = v.version_no
                break

    tags = [t.name for t in (d.tags or [])]
    if not tags:
        try:
            tags = [t.name for t in (d.tags or [])]
        except Exception:
            tags = []

    return DocumentOut(
        id=d.id,
        original_name=d.original_name,
        file_type=d.file_type,
        file_size=d.file_size,
        mime_type=d.mime_type,
        description=d.description,
        owner_id=d.owner_id,
        owner_name=d.owner.real_name if d.owner else None,
        matter_id=d.matter_id,
        matter_title=d.matter.title if d.matter else None,
        category_id=d.category_id,
        category_name=d.category.name if d.category else None,
        status=d.status,
        current_version_id=d.current_version_id,
        current_version_no=current_version_no,
        permission_scope=getattr(d, 'permission_scope', 'matter'),
        is_deleted=d.is_deleted,
        tags=tags,
        created_at=d.created_at.isoformat() if d.created_at else None,
        updated_at=d.updated_at.isoformat() if d.updated_at else None,
        extracted_text_length=len(d.extracted_text) if d.extracted_text else None,
    )


# ---------- Document Routes ----------

@router.get("/", response_model=DocumentListResponse)
async def list_documents(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=500),
    keyword: str = Query(None),
    file_type: str = Query(None),
    category_id: int = Query(None),
    matter_id: int = Query(None),
    tag: str = Query(None),
    status: str = Query(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    pagination = PaginationParams(page=page, page_size=page_size)

    # Resolve tag name to ID
    tag_id = None
    if tag:
        tag_result = await db.execute(select(Tag.id).where(Tag.name == tag))
        tag_id = tag_result.scalar_one_or_none()

    filters = {}
    if keyword:
        filters["keyword"] = keyword
    if file_type:
        filters["file_type"] = file_type
    if category_id:
        filters["category_id"] = category_id
    if matter_id:
        filters["matter_id"] = matter_id
    if status:
        filters["status"] = status
    if tag_id:
        filters["tag_id"] = tag_id

    current_user_dict = _user_to_dict(current_user)
    docs, total = await document_service.get_documents(
        db, pagination, filters, current_user_dict
    )

    doc_list = [_doc_to_out(d) for d in docs]
    total_pages = (total + page_size - 1) // page_size if total > 0 else 0
    return DocumentListResponse(
        items=doc_list, total=total, page=page, page_size=page_size, total_pages=total_pages,
    )


@router.post("/upload", response_model=DocumentOut)
async def upload_document(
    file: UploadFile = File(...),
    matter_id: int | None = Form(None),
    category_id: int | None = Form(None),
    description: str | None = Form(None),
    tag_ids: str | None = Form(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if not file.filename:
        raise ValidationException(detail="File name is required")

    if not is_allowed_file(file.filename):
        raise ValidationException(detail=f"File type not allowed: {file.filename}")

    content = await file.read()
    if not content:
        raise ValidationException(detail="Empty file")

    file_size = len(content)
    settings = __import__('app.config', fromlist=['get_settings']).get_settings()
    if file_size > settings.MAX_FILE_SIZE:
        raise ValidationException(detail=f"File too large, max {settings.MAX_FILE_SIZE // (1024*1024)}MB")

    file_type = detect_file_type(file.filename, file.content_type or "application/octet-stream")
    checksum = compute_sha256(content)
    storage_path = generate_storage_path(file.filename, matter_id)
    mime_type = file.content_type or "application/octet-stream"

    # Upload to MinIO
    minio_client.upload_file(storage_path, content, content_type=mime_type)

    # Resolve tags
    tag_list = []
    parsed_tag_ids = []
    if tag_ids:
        try:
            parsed_tag_ids = [int(t.strip()) for t in tag_ids.split(",") if t.strip()]
            if parsed_tag_ids:
                tag_result = await db.execute(select(Tag).where(Tag.id.in_(parsed_tag_ids)))
                tag_list = tag_result.scalars().all()
        except ValueError:
            pass

    doc = Document(
        original_name=file.filename,
        file_type=file_type,
        file_size=file_size,
        mime_type=mime_type,
        storage_path=storage_path,
        description=description,
        owner_id=current_user.id,
        matter_id=matter_id,
        category_id=category_id,
        status="draft",
    )
    db.add(doc)
    await db.flush()

    version = DocumentVersion(
        document_id=doc.id,
        version_no=1,
        file_path=storage_path,
        file_size=file_size,
        upload_user_id=current_user.id,
        change_note="Initial upload",
        is_official=False,
        checksum=checksum,
    )
    db.add(version)
    await db.flush()

    doc.current_version_id = version.id
    if tag_list:
        doc.tags = tag_list

    await db.commit()

    # Fire-and-forget text extraction for intelligence pipeline
    doc_id_for_extract = doc.id
    file_type_for_extract = file_type
    original_name_for_extract = file.filename
    asyncio.create_task(_background_extract(db, doc_id_for_extract, content, file_type_for_extract, original_name_for_extract))

    # Re-query with eager loads
    result = await db.execute(
        select(Document)
        .options(
            selectinload(Document.owner),
            selectinload(Document.matter),
            selectinload(Document.category),
            selectinload(Document.tags),
        )
        .where(Document.id == doc.id)
    )
    doc = result.scalar_one()

    return DocumentOut(
        id=doc.id,
        original_name=doc.original_name,
        file_type=doc.file_type,
        file_size=doc.file_size,
        mime_type=doc.mime_type,
        description=doc.description,
        owner_id=doc.owner_id,
        owner_name=doc.owner.real_name if doc.owner else None,
        matter_id=doc.matter_id,
        matter_title=doc.matter.title if doc.matter else None,
        category_id=doc.category_id,
        category_name=doc.category.name if doc.category else None,
        status=doc.status,
        current_version_id=doc.current_version_id,
        current_version_no=1,
        permission_scope=doc.permission_scope,
        is_deleted=doc.is_deleted,
        tags=[t.name for t in (doc.tags or [])],
        created_at=doc.created_at.isoformat() if doc.created_at else None,
        updated_at=doc.updated_at.isoformat() if doc.updated_at else None,
    )


@router.post("/upload/init", response_model=PendingUpload)
async def init_chunked_upload(
    filename: str = Form(...),
    total_chunks: int = Form(...),
    file_size: int = Form(...),
    current_user: User = Depends(get_current_user),
):
    import uuid
    upload_id = uuid.uuid4().hex
    return PendingUpload(
        id=upload_id, filename=filename,
        uploaded_chunks=0, total_chunks=total_chunks, is_complete=False,
    )


@router.post("/upload/chunk")
async def upload_chunk(
    upload_id: str = Form(...),
    chunk_index: int = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    import os, tempfile
    tmp_dir = os.path.join(tempfile.gettempdir(), "odms_chunks", upload_id)
    os.makedirs(tmp_dir, exist_ok=True)
    chunk_path = os.path.join(tmp_dir, f"chunk_{chunk_index:06d}")
    content = await file.read()
    with open(chunk_path, "wb") as f:
        f.write(content)
    return {"upload_id": upload_id, "chunk_index": chunk_index, "uploaded": True}


@router.post("/upload/complete", response_model=DocumentOut)
async def complete_chunked_upload(
    data: CompleteUploadRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    import os, tempfile, shutil

    tmp_dir = os.path.join(tempfile.gettempdir(), "odms_chunks", data.upload_id)
    if not os.path.isdir(tmp_dir):
        raise NotFoundException(resource="Upload session")

    chunks = sorted([f for f in os.listdir(tmp_dir) if f.startswith("chunk_")])
    merged_data = b""
    for chunk_file in chunks:
        with open(os.path.join(tmp_dir, chunk_file), "rb") as f:
            merged_data += f.read()

    if not merged_data:
        raise ValidationException(detail="No data received")

    file_type = detect_file_type(data.filename, "application/octet-stream")
    checksum = compute_sha256(merged_data)
    storage_path = generate_storage_path(data.filename, data.matter_id)
    mime_type = "application/octet-stream"

    minio_client.upload_file(storage_path, merged_data, content_type=mime_type)

    tag_list = []
    if data.tag_ids:
        tag_result = await db.execute(select(Tag).where(Tag.id.in_(data.tag_ids)))
        tag_list = tag_result.scalars().all()

    doc = Document(
        original_name=data.filename,
        file_type=file_type,
        file_size=len(merged_data),
        mime_type=mime_type,
        storage_path=storage_path,
        description=data.description,
        owner_id=current_user.id,
        matter_id=data.matter_id,
        category_id=data.category_id,
        status="draft",
    )
    db.add(doc)
    await db.flush()

    version = DocumentVersion(
        document_id=doc.id,
        version_no=1,
        file_path=storage_path,
        file_size=len(merged_data),
        upload_user_id=current_user.id,
        change_note="Initial upload (chunked)",
        is_official=False,
        checksum=checksum,
    )
    db.add(version)
    await db.flush()

    doc.current_version_id = version.id
    if tag_list:
        doc.tags = tag_list

    await db.commit()

    # Fire-and-forget text extraction
    asyncio.create_task(_background_extract(db, doc.id, merged_data, file_type, data.filename))

    # Dispatch preview PDF generation (LibreOffice, Celery with independent session)
    from app.tasks.preview_tasks import convert_to_pdf
    convert_to_pdf.delay(doc.id, storage_path, file_type)

    # Re-query with eager loads
    result = await db.execute(
        select(Document)
        .options(
            selectinload(Document.owner),
            selectinload(Document.matter),
            selectinload(Document.category),
            selectinload(Document.tags),
        )
        .where(Document.id == doc.id)
    )
    doc = result.scalar_one()

    try:
        shutil.rmtree(tmp_dir)
    except Exception:
        pass

    return DocumentOut(
        id=doc.id,
        original_name=doc.original_name,
        file_type=doc.file_type,
        file_size=doc.file_size,
        mime_type=doc.mime_type,
        description=doc.description,
        owner_id=doc.owner_id,
        owner_name=doc.owner.real_name if doc.owner else None,
        matter_id=doc.matter_id,
        matter_title=doc.matter.title if doc.matter else None,
        category_id=doc.category_id,
        category_name=doc.category.name if doc.category else None,
        status=doc.status,
        current_version_id=doc.current_version_id,
        current_version_no=1,
        permission_scope=doc.permission_scope,
        is_deleted=doc.is_deleted,
        tags=[t.name for t in (doc.tags or [])],
        created_at=doc.created_at.isoformat() if doc.created_at else None,
        updated_at=doc.updated_at.isoformat() if doc.updated_at else None,
    )


# ---------- Category Routes ----------

@router.get("/categories", response_model=list[CategoryOut])
async def list_categories(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    categories = await document_service.get_document_categories(db)
    return [CategoryOut.model_validate(c) for c in categories]


@router.post("/categories", response_model=CategoryOut)
async def create_category(
    data: CategoryCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    cat = await document_service.create_category(db, data)
    return CategoryOut.model_validate(cat)


@router.put("/categories/{cat_id}", response_model=CategoryOut)
async def update_category(
    cat_id: int,
    data: CategoryUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    cat = await document_service.update_category(db, cat_id, data)
    return CategoryOut.model_validate(cat)


@router.delete("/categories/{cat_id}")
async def delete_category(
    cat_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await document_service.delete_category(db, cat_id)
    return {"detail": "Category deleted"}


# ---------- Tag Routes ----------

@router.get("/tags", response_model=list[TagOut])
async def list_tags(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    tags = await document_service.get_tags(db)
    return [TagOut.model_validate(t) for t in tags]


@router.post("/tags", response_model=TagOut)
async def create_tag(
    data: TagCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    tag = await document_service.create_tag(db, data)
    return TagOut.model_validate(tag)


@router.delete("/tags/{tag_id}")
async def delete_tag(
    tag_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await document_service.delete_tag(db, tag_id)
    return {"detail": "Tag deleted"}


@router.get("/{doc_id}", response_model=DocumentOut)
async def get_document(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)
    # Additional access check via the route helper
    await _check_document_access(current_user, doc, db)
    return _doc_to_out(doc)


@router.put("/{doc_id}", response_model=DocumentOut)
async def update_document(
    doc_id: int,
    data: DocumentUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.update_document(db, doc_id, data, current_user_dict)
    # Re-query with eager loads for response
    doc = await document_service.get_document(db, doc_id, current_user_dict)
    return _doc_to_out(doc)


@router.delete("/{doc_id}")
async def delete_document(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete a document. Owner, admin, or dept leader can delete."""
    current_user_dict = _user_to_dict(current_user)
    await document_service.soft_delete_document(db, doc_id, current_user_dict)
    return {"detail": "Document deleted successfully"}


@router.get("/{doc_id}/download")
async def download_document(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)

    data = minio_client.download_file(doc.storage_path)
    if data is None:
        raise NotFoundException(resource="File data")

    return StreamingResponse(
        BytesIO(data),
        media_type=doc.mime_type or "application/octet-stream",
        headers={
            "Content-Disposition": f'attachment; filename="{doc.original_name}"',
        },
    )


@router.get("/{doc_id}/preview")
async def get_preview_url(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)

    preview_pdf_path = getattr(doc, 'preview_pdf_path', None)
    if preview_pdf_path:
        url = minio_client.get_presigned_url(preview_pdf_path)
        return {"url": url, "status": "ready"}
    else:
        url = minio_client.get_presigned_url(doc.storage_path)
        return {"url": url, "status": "ready", "is_original": True}


@router.get("/{doc_id}/versions", response_model=list[VersionOut])
async def list_versions(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)

    versions = await version_service.get_versions(db, doc_id)
    return [
        VersionOut(
            id=v.id, version_no=v.version_no, file_size=v.file_size,
            upload_user_name=v.upload_user.real_name if v.upload_user else None,
            change_note=v.change_note, is_official=v.is_official,
            checksum=v.checksum,
            created_at=v.created_at.isoformat() if v.created_at else None,
        )
        for v in versions
    ]


@router.post("/{doc_id}/versions", response_model=VersionOut)
async def upload_new_version(
    doc_id: int,
    file: UploadFile = File(...),
    change_note: str | None = Form(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)

    if not file.filename:
        raise ValidationException(detail="File name is required")

    content = await file.read()
    if not content:
        raise ValidationException(detail="Empty file")

    settings = __import__('app.config', fromlist=['get_settings']).get_settings()
    if len(content) > settings.MAX_FILE_SIZE:
        raise ValidationException(detail="File too large")

    storage_path = generate_storage_path(file.filename, doc.matter_id)
    checksum = compute_sha256(content)

    minio_client.upload_file(
        storage_path, content,
        content_type=file.content_type or "application/octet-stream",
    )

    file_info = {
        "storage_path": storage_path,
        "file_size": len(content),
        "file_type": detect_file_type(file.filename, file.content_type or "application/octet-stream"),
        "mime_type": file.content_type or "application/octet-stream",
        "checksum": checksum,
    }

    version = await version_service.upload_new_version(
        db, doc_id, file_info, current_user.id, change_note, set_as_official=True
    )

    return VersionOut(
        id=version.id, version_no=version.version_no,
        file_size=version.file_size,
        upload_user_name=current_user.real_name,
        change_note=version.change_note,
        is_official=version.is_official,
        checksum=version.checksum,
        created_at=version.created_at.isoformat() if version.created_at else None,
    )


@router.put("/{doc_id}/versions/{vid}/official")
async def set_version_official(
    doc_id: int,
    vid: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    _: None = Depends(check_permission(Permission.DOCUMENT_SET_OFFICIAL)),
):
    current_user_dict = _user_to_dict(current_user)
    version = await version_service.set_official_version(db, doc_id, vid, current_user_dict)
    return {"detail": "Official version set", "version_no": version.version_no}


@router.post("/{doc_id}/lock")
async def lock_document(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
    _: None = Depends(check_permission(Permission.DOCUMENT_LOCK)),
):
    current_user_dict = _user_to_dict(current_user)
    # Verify document exists and user has access
    await document_service.get_document(db, doc_id, current_user_dict)

    lock = await lock_service.lock_document(db, doc_id, current_user.id)
    return {"detail": "Document locked", "expires_at": lock.expires_at.isoformat()}


@router.delete("/{doc_id}/lock")
async def unlock_document(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    # Verify document exists and user has access
    await document_service.get_document(db, doc_id, current_user_dict)

    await lock_service.unlock_document(db, doc_id, current_user.id)
    return {"detail": "Document unlocked"}


@router.get("/{doc_id}/lock-status", response_model=LockStatusOut)
async def get_lock_status(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    status_data = await lock_service.get_lock_status(db, doc_id)
    if status_data["is_locked"]:
        locker = status_data["locked_by"]
        return LockStatusOut(
            is_locked=True,
            locked_by=locker.get("id") if isinstance(locker, dict) else locker,
            locked_by_name=locker.get("real_name") if isinstance(locker, dict) else None,
            locked_at=status_data.get("locked_at"),
            expires_at=status_data.get("expires_at"),
        )
    return LockStatusOut(
        is_locked=False, locked_by=None, locked_by_name=None,
        locked_at=None, expires_at=None,
    )


@router.post("/{doc_id}/reference", response_model=ReferenceOut)
async def add_cross_reference(
    doc_id: int,
    data: ReferenceCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    # Verify document exists and user has access
    await document_service.get_document(db, doc_id, current_user_dict)

    ref = await reference_service.add_reference(
        db, doc_id, data.matter_id, current_user.id, data.is_readonly
    )

    # Get matter title
    matter_result = await db.execute(select(Matter).where(Matter.id == data.matter_id))
    matter = matter_result.scalar_one_or_none()

    return ReferenceOut(
        id=ref.id, matter_id=ref.matter_id,
        matter_title=matter.title if matter else None,
        is_readonly=ref.is_readonly,
        added_by_name=current_user.real_name,
        created_at=ref.created_at.isoformat() if ref.created_at else None,
    )


@router.get("/{doc_id}/references", response_model=list[ReferenceOut])
async def list_references(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    # Verify document exists and user has access
    await document_service.get_document(db, doc_id, current_user_dict)

    refs = await reference_service.get_references(db, doc_id)
    return [
        ReferenceOut(
            id=r.id, matter_id=r.matter_id,
            matter_title=r.matter.title if r.matter else None,
            is_readonly=r.is_readonly,
            added_by_name=r.adder.real_name if r.adder else None,
            created_at=r.created_at.isoformat() if r.created_at else None,
        )
        for r in refs
    ]


@router.delete("/{doc_id}/references/{ref_id}")
async def remove_reference(
    doc_id: int,
    ref_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    current_user_dict = _user_to_dict(current_user)
    # Verify document exists and user has access
    await document_service.get_document(db, doc_id, current_user_dict)

    await reference_service.remove_reference(db, ref_id, current_user.id)
    return {"detail": "Reference removed"}


# ---------- Approval Schemas ----------

class SubmitForReviewRequest(BaseModel):
    reviewer_ids: list[int]

class ReviewActionRequest(BaseModel):
    comment: str | None = None

class ReviewOut(BaseModel):
    id: int
    document_id: int
    review_level: int
    reviewer_id: int
    reviewer_name: str | None
    status: str
    comment: str | None
    reviewed_at: str | None
    created_at: str | None

    class Config:
        from_attributes = True


# ---------- Approval Routes ----------

@router.post("/{doc_id}/submit-review", response_model=DocumentOut)
async def submit_for_review(
    doc_id: int,
    data: SubmitForReviewRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Submit document for approval with a chain of reviewers."""
    doc = await review_service.submit_for_review(
        db, doc_id, data.reviewer_ids, current_user.id
    )
    # Re-query for full response
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)
    return _doc_to_out(doc)


@router.post("/{doc_id}/approve/{review_level}", response_model=ReviewOut)
async def approve_document(
    doc_id: int,
    review_level: int,
    data: ReviewActionRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Approve document at a specific review level."""
    review = await review_service.approve_document(
        db, doc_id, review_level, current_user.id, data.comment
    )
    review_result = await db.execute(
        select(DocumentReview)
        .options(selectinload(DocumentReview.reviewer))
        .where(DocumentReview.id == review.id)
    )
    review = review_result.scalar_one()
    return ReviewOut(
        id=review.id,
        document_id=review.document_id,
        review_level=review.review_level,
        reviewer_id=review.reviewer_id,
        reviewer_name=review.reviewer.real_name if review.reviewer else None,
        status=review.status,
        comment=review.comment,
        reviewed_at=review.reviewed_at.isoformat() if review.reviewed_at else None,
        created_at=review.created_at.isoformat() if review.created_at else None,
    )


@router.post("/{doc_id}/reject/{review_level}", response_model=ReviewOut)
async def reject_document(
    doc_id: int,
    review_level: int,
    data: ReviewActionRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Reject document at a specific review level."""
    review = await review_service.reject_document(
        db, doc_id, review_level, current_user.id, data.comment
    )
    review_result = await db.execute(
        select(DocumentReview)
        .options(selectinload(DocumentReview.reviewer))
        .where(DocumentReview.id == review.id)
    )
    review = review_result.scalar_one()
    return ReviewOut(
        id=review.id,
        document_id=review.document_id,
        review_level=review.review_level,
        reviewer_id=review.reviewer_id,
        reviewer_name=review.reviewer.real_name if review.reviewer else None,
        status=review.status,
        comment=review.comment,
        reviewed_at=review.reviewed_at.isoformat() if review.reviewed_at else None,
        created_at=review.created_at.isoformat() if review.created_at else None,
    )


@router.get("/{doc_id}/reviews", response_model=list[ReviewOut])
async def get_reviews(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get all reviews for a document."""
    reviews = await review_service.get_reviews(db, doc_id)
    return [
        ReviewOut(
            id=r.id,
            document_id=r.document_id,
            review_level=r.review_level,
            reviewer_id=r.reviewer_id,
            reviewer_name=r.reviewer.real_name if r.reviewer else None,
            status=r.status,
            comment=r.comment,
            reviewed_at=r.reviewed_at.isoformat() if r.reviewed_at else None,
            created_at=r.created_at.isoformat() if r.created_at else None,
        )
        for r in reviews
    ]


@router.get("/pending-reviews/mine")
async def my_pending_reviews(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get all pending reviews assigned to the current user."""
    reviews = await review_service.get_pending_reviews_for_user(db, current_user.id)
    result = []
    for r in reviews:
        doc = r.document
        result.append({
            "id": r.id,
            "document_id": r.document_id,
            "document_name": doc.original_name if doc else None,
            "document_owner": doc.owner.real_name if doc and doc.owner else None,
            "review_level": r.review_level,
            "status": r.status,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        })
    return result


# ---------- Export Routes ----------

@router.get("/export/excel")
async def export_documents_excel(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    pagination = PaginationParams(page=1, page_size=10000)
    filters = {}
    current_user_dict = _user_to_dict(current_user)
    docs, _ = await document_service.get_documents(db, pagination, filters, current_user_dict)

    wb = Workbook()
    ws = wb.active
    ws.title = "\u6587\u6863\u5217\u8868"
    ws.append(["ID", "\u6587\u4ef6\u540d", "\u7c7b\u578b", "\u5927\u5c0f(KB)", "\u4e0a\u4f20\u8005", "\u72b6\u6001", "\u4e0a\u4f20\u65f6\u95f4"])

    for d in docs:
        ws.append([
            d.id, d.original_name, d.file_type, round(d.file_size / 1024, 1),
            d.owner.real_name if d.owner else '', d.status,
            d.created_at.strftime('%Y-%m-%d %H:%M') if d.created_at else ''
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=documents.xlsx"}
    )


# ---------- Document Intelligence Routes ----------

from app.services.document_intelligence import (
    suggest_category, suggest_tags, find_similar_documents, find_similar_by_vector,
    find_similar_by_text, extract_and_store_text, get_document_graph_data,
)


class SimilarDocumentOut(BaseModel):
    document_id: int
    original_name: str
    file_type: str
    description: str | None
    status: str
    similarity_score: float
    headline: str | None


class CategorySuggestionOut(BaseModel):
    category_id: int
    category_name: str
    score: int


class TagSuggestionOut(BaseModel):
    tag_id: int
    tag_name: str
    matched_keyword: str


class GraphDataOut(BaseModel):
    nodes: list[dict]
    links: list[dict]
    categories: list[dict]


@router.post("/{doc_id}/extract-text")
async def trigger_text_extraction(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Extract text from an existing document and store it for intelligence features."""
    current_user_dict = _user_to_dict(current_user)
    doc = await document_service.get_document(db, doc_id, current_user_dict)

    if not is_supported(doc.file_type):
        return {"detail": f"Text extraction not supported for file type: {doc.file_type}", "extracted": False}

    try:
        data = minio_client.download_file(doc.storage_path)
        if not data:
            raise NotFoundException(resource="File data")
        extracted = await extract_and_store_text(db, doc_id, data, doc.file_type, doc.original_name)
        if extracted:
            return {"detail": "Text extracted successfully", "extracted": True, "length": len(extracted)}
        else:
            return {"detail": "No text could be extracted from this file", "extracted": False}
    except Exception as e:
        return {"detail": f"Extraction failed: {str(e)}", "extracted": False}


@router.get("/{doc_id}/similar", response_model=list[SimilarDocumentOut])
async def get_similar_documents(
    doc_id: int,
    limit: int = Query(10, ge=1, le=50),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Find documents similar to the given document."""
    current_user_dict = _user_to_dict(current_user)
    await document_service.get_document(db, doc_id, current_user_dict)  # Verify access
    results = await find_similar_documents(db, doc_id, limit)
    return [SimilarDocumentOut(**r) for r in results]


@router.get("/{doc_id}/similar-vector", response_model=list[SimilarDocumentOut])
async def get_similar_documents_vector(
    doc_id: int,
    limit: int = Query(10, ge=1, le=50),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Find similar documents using vector similarity (pgvector cosine distance)."""
    current_user_dict = _user_to_dict(current_user)
    await document_service.get_document(db, doc_id, current_user_dict)
    results = await find_similar_by_vector(db, doc_id, limit)
    return [SimilarDocumentOut(**r) for r in results]


@router.get("/{doc_id}/graph", response_model=GraphDataOut)
async def get_document_graph(
    doc_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get force-directed graph data for document relationships."""
    current_user_dict = _user_to_dict(current_user)
    await document_service.get_document(db, doc_id, current_user_dict)  # Verify access
    data = await get_document_graph_data(db, doc_id)
    return GraphDataOut(**data)


@router.get("/intelligence/suggest-category", response_model=list[CategorySuggestionOut])
async def suggest_document_category(
    file_name: str = Query(""),
    text: str = Query(""),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Suggest categories based on file name and/or text content."""
    results = await suggest_category(db, file_name, text or None)
    return [CategorySuggestionOut(**r) for r in results]


@router.get("/intelligence/suggest-tags", response_model=list[TagSuggestionOut])
async def suggest_document_tags(
    text: str = Query(""),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Suggest tags based on text content."""
    results = await suggest_tags(db, text or None)
    return [TagSuggestionOut(**r) for r in results]


@router.get("/intelligence/search-similar", response_model=list[SimilarDocumentOut])
async def search_similar_documents(
    text: str = Query(...),
    exclude_id: int = Query(None),
    limit: int = Query(10, ge=1, le=50),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Search for documents similar to arbitrary text (for upload-time duplicate detection)."""
    results = await find_similar_by_text(db, text, exclude_id, limit)
    return [SimilarDocumentOut(**r) for r in results]
